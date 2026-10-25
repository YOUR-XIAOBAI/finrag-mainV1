import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

BASE_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE_DIR))

from custom_loader import WindowsSafeDirectoryLoader
from document_processing import load_and_process_documents
from rag_pipeline import FinlandTravelRAG


class RetrievalTests(unittest.TestCase):
    def test_original_samples_and_chinese_queries(self):
        rag = FinlandTravelRAG()
        self.assertEqual(rag.source_count, 4)
        self.assertEqual(len(rag.warnings), 2)
        for query, filename in [
            ("芬兰什么时候看极光？", "northern_lights_guide.pdf"),
            ("赫尔辛基的交通", "visitng_helsinki.txt"),
            ("夏天要带什么", "summer_in_finland.txt"),
            ("Finnish sauna etiquette", "finnish_culture.docx"),
        ]:
            with self.subTest(query=query):
                self.assertTrue(rag._retrieve(query, 1)[0].metadata["source"].endswith(filename))
        self.assertIn("September to March", rag.generate_response("看极光", k=1))
        self.assertIn("未找到", rag.generate_response("quantum superconductivity"))
        self.assertIn("未找到", rag.generate_response("日本签证怎么办"))
        self.assertIn("请输入", rag.generate_response("  "))
        self.assertEqual(rag._retrieve("aurora", 0), [])

    def test_real_word_pdf_and_corrupt_file(self):
        from docx import Document
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

        with tempfile.TemporaryDirectory(prefix="芬兰资料-") as folder:
            root = Path(folder)
            word = Document()
            word.add_paragraph("Helsinki test paragraph")
            word.add_table(rows=1, cols=1).cell(0, 0).text = "Table transport"
            word.save(root / "真实.docx")
            writer = PdfWriter()
            page = writer.add_blank_page(width=300, height=300)
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                     NameObject("/Subtype"): NameObject("/Type1"),
                                     NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
            stream = DecodedStreamObject()
            stream.set_data(b"BT /F1 12 Tf 30 200 Td (Aurora in Lapland) Tj ET")
            page[NameObject("/Contents")] = writer._add_object(stream)
            writer.write(root / "真实.pdf")
            (root / "broken.pdf").write_bytes(b"%PDF-invalid\x00")
            (root / "empty.txt").write_text("", encoding="utf-8")
            (root / "ignored.exe").write_bytes(b"binary")
            loader = WindowsSafeDirectoryLoader(root)
            docs = loader.load()
            self.assertEqual(len(docs), 2)
            self.assertEqual(len(loader.warnings), 2)
            self.assertTrue(any("Table transport" in d.page_content for d in docs))
            self.assertTrue(any("Aurora in Lapland" in d.page_content and d.metadata["page"] == 1 for d in docs))

    def test_chunking_and_empty_directory(self):
        with tempfile.TemporaryDirectory() as folder:
            with self.assertRaisesRegex(ValueError, "没有读到"):
                load_and_process_documents(folder)
            (Path(folder) / "long.txt").write_text("Aurora Lapland\n" * 100, encoding="utf-8")
            chunks = load_and_process_documents(folder, chunk_size=100, chunk_overlap=20)
            self.assertGreater(len(chunks), 1)
            self.assertTrue(all(0 < len(d.page_content) <= 100 for d in chunks))
            with self.assertRaises(ValueError):
                load_and_process_documents(folder, chunk_size=100, chunk_overlap=100)

    def test_optional_ai_error_is_actionable(self):
        rag = FinlandTravelRAG(mode="ai")
        with patch("rag_pipeline.ai_dependencies_available", return_value=False):
            with self.assertRaisesRegex(RuntimeError, "install_ai.bat"):
                rag.generate_response("aurora")


class AppTests(unittest.TestCase):
    def test_chat_and_rebuild(self):
        from streamlit.testing.v1 import AppTest

        app = AppTest.from_file(str(BASE_DIR / "app.py")).run(timeout=20)
        app.selectbox[0].select("本地资料检索").run(timeout=20)
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(app.metric[0].value, "4 个文件 / 4 个片段")
        app.chat_input[0].set_value("芬兰什么时候看极光？").run(timeout=20)
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(len(app.chat_message), 2)
        self.assertIn("September to March", app.chat_message[1].markdown[0].value)
        app.button[0].click().run(timeout=20)
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(len(app.chat_message), 0)


if __name__ == "__main__":
    unittest.main()

"""Local loading, including the original text-only PDF/DOCX samples."""
from __future__ import annotations

import logging
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

logger = logging.getLogger(__name__)
SUPPORTED_SUFFIXES = {".txt", ".md", ".pdf", ".docx"}


@dataclass
class Document:
    page_content: str
    metadata: dict = field(default_factory=dict)


class WindowsSafeDirectoryLoader:
    def __init__(self, root_dir: str | Path, max_workers: int = 8) -> None:
        self.root_dir = Path(root_dir).resolve()
        if not self.root_dir.is_dir():
            raise FileNotFoundError(f"Documents folder does not exist: {self.root_dir}")
        self.warnings: list[str] = []

    @staticmethod
    def _read_text(path: Path) -> str:
        raw = path.read_bytes()
        if b"\x00" in raw:
            raise ValueError("The file contains binary data and cannot be read as text")
        for encoding in ("utf-8-sig", "gb18030", "cp1252"):
            try:
                return raw.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise ValueError("Unrecognized text encoding. Please convert the file to UTF-8")

    def _load_single(self, path: Path) -> list[Document]:
        source = path.relative_to(self.root_dir).as_posix()
        suffix = path.suffix.lower()
        if suffix == ".pdf" and path.read_bytes()[:5] == b"%PDF-":
            from pypdf import PdfReader

            docs = [Document(page.extract_text() or "", {"source": source, "page": i + 1})
                    for i, page in enumerate(PdfReader(str(path)).pages)]
            if not any(doc.page_content.strip() for doc in docs):
                raise ValueError("The PDF has no extractable text. Scanned documents need OCR first")
            return docs
        if suffix == ".docx" and zipfile.is_zipfile(path):
            from docx import Document as WordDocument

            word = WordDocument(str(path))
            blocks = [p.text for p in word.paragraphs]
            blocks += [" | ".join(cell.text for cell in row.cells)
                       for table in word.tables for row in table.rows]
            text = "\n".join(blocks)
        else:
            text = self._read_text(path)
            if suffix in {".pdf", ".docx"}:
                self.warnings.append(f"{source} contains plain text despite its extension. Loaded as text.")
        if not text.strip():
            raise ValueError("The file is empty")
        return [Document(text.strip(), {"source": source})]

    def load(self) -> list[Document]:
        self.warnings.clear()
        docs: list[Document] = []
        for path in sorted(self.root_dir.rglob("*")):
            if not path.is_file() or path.suffix.lower() not in SUPPORTED_SUFFIXES:
                continue
            if path.name.startswith((".", "~$")):
                continue
            try:
                docs.extend(self._load_single(path))
            except Exception as exc:
                message = f"Unable to read {path.name}: {exc}"
                self.warnings.append(message)
                logger.warning(message)
        return docs
